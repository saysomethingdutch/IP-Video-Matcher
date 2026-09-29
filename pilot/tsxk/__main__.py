"""吞噬星空 Q–AC (WO-102 §3–4).

    python -m pilot.tsxk run cases/tsxk.yaml [--rows 7813,7814] [--no-local] [--workers 4]
    python -m pilot.tsxk samples cases/tsxk.yaml [--videos 13,16]      # the example recordings, for checking
    python -m pilot.tsxk check-list cases/tsxk.yaml                    # the 30-row list (input columns only)
    python -m pilot.tsxk check-score cases/tsxk.yaml --owner-approved  # reads the human columns: Owner only
    python -m pilot.tsxk usage cases/tsxk.yaml [--since 2026-09-29T18:00]  # tokens and time per kind of call

Needs the timelines first (python -m pilot.timeline …). Cloud calls: launch with zsh -lic '…'.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

from pilot.semantic import llm
from pilot.timeline.cloud import share_slots

from . import check as CK
from . import usage as US
from .config import TsxkCase
from .pipeline import process_row, save, system_prompt
from .samples import check_rows

OUT_COLUMNS = ["D", "F", "L", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC"]


def _load(path: Path) -> dict | None:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def _csv(results: dict[str, dict], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["序号", *OUT_COLUMNS, "标黄", "标红"])
        for rid, r in results.items():
            fl = r["flags"]
            w.writerow([rid, *[(r["fields"].get(c) or {}).get("value") for c in OUT_COLUMNS],
                        "；".join(x["reason"] for x in fl if x["level"] == "yellow"),
                        "；".join(x["reason"] for x in fl if x["level"] == "red")])


def sample_inputs(case: TsxkCase) -> dict[str, dict]:
    """Input columns of the guidance examples that match the sample recordings (by 序号)."""
    g = case.raw["guidance"]
    keep = set(case.raw["workbook"]["input_columns"])
    ws = openpyxl.load_workbook(Path(g["dir"]) / g["examples"], read_only=True, data_only=True).worksheets[0]
    by_sid = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        d = {openpyxl.utils.get_column_letter(i + 1): v for i, v in enumerate(values)}
        if d.get("B") is not None:
            by_sid[int(d["B"])] = {k: v for k, v in d.items() if k in keep}
    return {name: by_sid.get(int(sid), {}) for name, sid in (g.get("sample_rows") or {}).items()}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m pilot.tsxk")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("run", "samples", "check-list", "check-run", "check-freeze", "check-score", "usage"):
        p = sub.add_parser(name)
        p.add_argument("config")
        p.add_argument("--rows", default="")
        p.add_argument("--videos", default="13,14,15,16,17")
        p.add_argument("--no-local", action="store_true")
        p.add_argument("--workers", type=int, default=4)
        p.add_argument("--owner-approved", action="store_true")
        p.add_argument("--machine", default="")
        p.add_argument("--set", default="a", choices=("a", "b"))
        p.add_argument("--since", default="")
        p.add_argument("--until", default="")
        p.add_argument("--run", default="")
    args = ap.parse_args(argv)
    llm.LOCAL_ALLOWED = False                      # GLM only (D-037); never the local model (Codex's Ollama)
    case = TsxkCase(args.config)
    share_slots(case.case_dir / ".cloud_slots")     # one 8/4 limit per provider across processes (D-058)
    llm.CALL_LOG = case.case_dir / "usage" / "calls.jsonl"     # every cloud request with its tokens (D-069)
    llm.CALL_RUN = " ".join(["tsxk", args.cmd] + ([f"--set {args.set}"] if args.cmd.startswith("check") else []))
    if args.cmd == "usage":
        print(US.markdown(US.read(llm.CALL_LOG, args.since, args.until, args.run)))
        return 0
    cs = case.check_set(args.set)
    if args.cmd == "check-list":
        if Path(cs["list"]).exists():                  # a drawn list is the Owner's record: never redraw over it
            print(f"名单已存在，不覆盖：{cs['list']}", file=sys.stderr)
            return 2
        rows = CK.filled_rows(Path(case.raw["workbook"]["path"]), cs["sheet"], cs["filled_column"])
        picked = CK.draw(rows, int(cs["size"]), int(cs["seed"]))
        CK.write_list(picked, case.path_of(cs["list"]))
        print(f"{len(rows)} 行已填，抽 {len(picked)} 行：", dict(__import__("collections").Counter(r["层"] for r in picked)))
        return 0
    if args.cmd == "check-freeze":
        return freeze(case, cs)
    if args.cmd == "check-score":
        if not args.owner_approved:
            print("检验会读取人工结论，需要 Owner 放行：加 --owner-approved", file=sys.stderr)
            return 2
        ids = check_rows(cs["list"])
        frozen = Path(args.machine) if args.machine else latest_frozen(cs)
        if frozen is None:
            print("没有冻结的机器结果：先运行 check-freeze", file=sys.stderr)
            return 2
        rows = {rid: r for rid in ids if (r := _load(frozen / "rows" / f"{rid}.json"))}
        machine = {rid: {c: (v or {}).get("value") for c, v in r["fields"].items()} for rid, r in rows.items()}
        report = CK.score(machine, CK.human_values(Path(case.raw["workbook"]["path"]), cs["sheet"], ids))
        report["frozen"] = str(frozen)
        tls = {rid: (cs["out"] / "timeline" / rid, _load(frozen / "timeline" / f"{rid}.json")) for rid in rows}
        audio = {rid: case.case_dir / "audio" / f"{rid}.json" for rid in rows}
        items = CK.disagreements(report, rows, tls, audio, {str(k): v for k, v in (cs.get("known_unreliable") or {}).items()})
        out = cs["out"]
        (out / "check_report.json").write_text(json.dumps({**report, "disagreements": items}, ensure_ascii=False, indent=1),
                                               encoding="utf-8")
        (out / "check_report.md").write_text(CK.markdown(report), encoding="utf-8")
        (out / "disagreements.md").write_text(CK.disagreements_md(items), encoding="utf-8")
        print(CK.markdown(report))
        print(f"分歧 {len(items)} 条：{out / 'disagreements.md'}")
        return 0
    rules = case.rules()
    system, info = system_prompt(case)
    print(f"风格样例：示例表 {info['guidance_examples']} 条，同事 {info['colleague_examples']} 条"
          + ("（检验名单未生成，未取同事样例）" if info["colleague_skipped"] else ""))
    backend = case.raw["models"]["text"]
    if args.cmd == "check-run":
        ids = sorted(check_rows(cs["list"]), key=int)
        ids = [r.strip() for r in args.rows.split(",") if r.strip()] or ids
        inputs = case.read_inputs(sheet=cs["sheet"], rows=ids)             # input columns only
        jobs = [(r, inputs.get(r, {}), case.webledger(r), _load(cs["out"] / "timeline" / r / "timeline.json"),
                 cs["out"] / "rows" / f"{r}.json", case.audio(r)) for r in ids]
    elif args.cmd == "samples":
        inputs = sample_inputs(case)
        jobs = [(n, inputs.get(n, {}), case.webledger(f"samples/{n}"), _load(case.case_dir / "samples" / n / "timeline.json"),
                 case.case_dir / "samples" / n / "fields.json", case.audio(f"samples/{n}")) for n in args.videos.split(",")]
    else:
        rows = [r.strip() for r in args.rows.split(",") if r.strip()] or case.rows
        inputs = case.read_inputs(rows=rows)
        jobs = [(r, inputs.get(r, {}), case.webledger(r), _load(case.case_dir / "timeline" / r / "timeline.json"),
                 case.case_dir / "rows" / f"{r}.json", case.audio(r)) for r in rows]

    def one(job):
        rid, inp, led, tl, path, audio = job
        res = process_row(inp, led, tl, rules, system, backend, audio)
        save(res, path)
        return rid, res

    results = {}
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as ex:
        for rid, res in ex.map(one, jobs):
            results[rid] = res
            print(f"{rid}: Y={res['fields'].get('Y', {}).get('value')} U={res['fields']['U']['value']} "
                  f"V={res['fields']['V']['value']} Z={res['fields'].get('Z', {}).get('value')} "
                  f"{'标红' if any(f['level'] == 'red' for f in res['flags']) else ''}", flush=True)
    _csv(results, cs["out"] / "check_fields.csv" if args.cmd == "check-run" else
         case.case_dir / ("samples_fields.csv" if args.cmd == "samples" else "fields.csv"))
    return 0


def latest_frozen(cs: dict) -> Path | None:
    snaps = sorted((cs["out"] / "frozen").glob("*/manifest.json"))
    return snaps[-1].parent if snaps else None


def freeze(case: TsxkCase, cs: dict) -> int:
    """Copy the check set's machine rows and timelines into check/frozen/<time>/ with SHA-256 and the code commit,
    before any human value is read (WO-102-R1 §2.1)."""
    import hashlib
    import shutil
    import subprocess
    from datetime import datetime
    ids = sorted(check_rows(cs["list"]), key=int)
    src = cs["out"]
    missing = [r for r in ids if not (src / "rows" / f"{r}.json").exists() or not (src / "timeline" / r / "timeline.json").exists()]
    if missing:
        print(f"还有 {len(missing)} 行没有机器结果，不能冻结：{', '.join(missing)}", file=sys.stderr)
        return 2
    dest = src / "frozen" / datetime.now().strftime("%Y%m%d-%H%M%S")
    files = {}
    for r in ids:
        for kind, path in (("rows", src / "rows" / f"{r}.json"), ("timeline", src / "timeline" / r / "timeline.json")):
            (dest / kind).mkdir(parents=True, exist_ok=True)
            target = dest / kind / f"{r}.json"
            shutil.copy2(path, target)
            files[f"{kind}/{r}.json"] = hashlib.sha256(target.read_bytes()).hexdigest()
    here = Path(__file__).parent
    commit = subprocess.run(["git", "rev-parse", "HEAD"], capture_output=True, text=True, cwd=here).stdout.strip()
    dirty = bool(subprocess.run(["git", "status", "--porcelain", "--", "pilot", "cases"], capture_output=True, text=True,
                                cwd=here).stdout.strip())
    (dest / "manifest.json").write_text(json.dumps({"frozen_at": datetime.now().isoformat(timespec="seconds"), "rows": len(ids),
                                                    "code_commit": commit, "code_dirty": dirty, "files": files},
                                                   ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"已冻结 {len(ids)} 行：{dest}（代码 {commit[:7]}{'，有未提交改动' if dirty else ''}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
