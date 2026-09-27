"""Read the two plaintiff workbooks without modifying them."""
from __future__ import annotations

from datetime import datetime, time
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET


def duration_seconds(value, *, qsjx=False):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        # Excel durations above 24 encoded minutes become a dated value.
        days = (value.date() - datetime(1899, 12, 31).date()).days
        if qsjx:
            return days * 1440 + value.hour * 60 + value.minute + value.second / 60
        return days * 86400 + value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, time):
        if qsjx:  # Excel stored m:ss as time(h=m, m=s).
            return value.hour * 60 + value.minute + value.second / 60
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)):
        return round(float(value) * 86400) if not qsjx else round(float(value) * 86400)
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        return duration_seconds(float(text), qsjx=qsjx)
    bits = text.split(":")
    if len(bits) == 3:
        return int(bits[0]) * 3600 + int(bits[1]) * 60 + float(bits[2])
    if len(bits) == 2:
        return int(bits[0]) * 60 + float(bits[1])
    return None


def read_qsjx(path: Path):
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True)["比对表"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    rows = []
    for excel_row, values in enumerate(iterator, 2):
        record = dict(zip(headers, values))
        name = str(record.get("录像名称") or "").split(".")[0]
        date = record.get("取证日期")
        if name not in {f"MVI_{n:04d}" for n in range(1, 9)}:
            continue
        if not isinstance(date, datetime) or date.date().isoformat() != "2023-10-12":
            continue
        rows.append({
            "row_id": str(record["总编号"]), "video": name,
            "bv": str(record.get("BV号") or ""),
            "title": str(record.get("播放/下载视频或操作信息") or ""),
            "uploader": str(record.get("上传者") or ""),
            "duration": duration_seconds(record.get("时长"), qsjx=True),
            "views_claimed": str(record.get("播放量") or ""),
            "part_claimed": str(record.get("侵权视频对应权利作品集数") or ""),
            "video_type": str(record.get("视频类型") or ""),
            "source_workbook": str(path), "source_sheet": "比对表", "source_row": excel_row,
        })
    if len(rows) != 169:
        raise ValueError(f"expected 169 七时吉祥 rows, got {len(rows)}")
    return rows


def read_configured(path: Path, spec: dict, videos: set[str]) -> list[dict]:
    """Read one batch using worksheet headers and row IDs supplied by the case."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[spec["sheet"]]
        header_max = 20 if spec.get("blind_holdout") else sheet.max_column
        headers = [str(value) if value is not None else "" for value in
                   next(sheet.iter_rows(min_row=1, max_row=1, max_col=header_max, values_only=True))]
        columns = spec["columns"]
        required = ("row_id", "video", "bv", "title", "uploader", "duration")
        missing = [key for key in required if not columns.get(key) or columns[key] not in headers]
        if missing:
            raise ValueError(f"workbook columns absent: {missing}")
        used_headers = [name for name in columns.values() if name]
        absent = [name for name in used_headers if name not in headers]
        if absent:
            raise ValueError(f"workbook columns absent: {absent}")
        # Read row values only through the last configured input field. In particular,
        # held-out human output columns must never enter the ledger pipeline.
        input_last_column = max(headers.index(name) + 1 for name in used_headers)
        if spec.get("blind_holdout") and input_last_column > 20:
            raise ValueError("blind holdout inputs must be within A–T")
        iterator = sheet.iter_rows(min_row=2, max_col=input_last_column, values_only=True)
        allowed = spec["row_ids"]
        ids = (set(map(str, allowed)) if isinstance(allowed, list)
               else {str(n) for n in range(int(allowed["first"]), int(allowed["last"]) + 1)})
        result = []
        for excel_row, values in enumerate(iterator, 2):
            record = dict(zip(headers, values))
            rid = str(record.get(columns["row_id"]) or "").removesuffix(".0")
            if rid not in ids:
                continue
            raw_video = str(record.get(columns["video"]) or "")
            video = Path(raw_video).stem
            # The workbook's recording name is a claim, not the search boundary.
            # Keep the row even when the claimed file is absent from this batch.
            date_key = columns.get("capture_date")
            if spec.get("capture_date") and date_key:
                date = record.get(date_key)
                actual = date.date().isoformat() if isinstance(date, datetime) else str(date or "")[:10]
                if actual != str(spec["capture_date"]):
                    continue

            def field(key: str) -> str:
                heading = columns.get(key)
                value = record.get(heading) if heading else None
                return str(value or "")

            result.append({
                "row_id": rid, "video": video, "bv": field("bv"), "title": field("title"),
                "uploader": field("uploader"),
                "duration": duration_seconds(record.get(columns["duration"]),
                                             qsjx=spec.get("duration_encoding") == "excel_mss"),
                "views_claimed": field("views_claimed"), "part_claimed": field("part_claimed"),
                "video_type": field("video_type"), "source_workbook": str(path),
                "source_sheet": spec["sheet"], "source_row": excel_row,
            })
        found = {r["row_id"] for r in result}
        if found != ids or len(result) != len(ids):
            raise ValueError(f"workbook batch mismatch: expected {len(ids)} rows, found {len(result)}; "
                             f"missing {sorted(ids - found)}")
        return result
    finally:
        workbook.close()


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _column_index(ref):
    index = 0
    for char in re.match(r"[A-Z]+", ref).group():
        index = index * 26 + ord(char) - 64
    return index - 1


def read_zhuyu(path: Path):
    with zipfile.ZipFile(path) as archive:
        shared_xml = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        shared = ["".join(t.text or "" for t in item.iter(NS + "t"))
                  for item in shared_xml.findall(NS + "si")]
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.get("Id"): item.get("Target") for item in relationships}
        sheet = next(s for s in workbook.find(NS + "sheets")
                     if s.get("name") == "取证日期排序（总表）")
        target = targets[sheet.get(REL + "id")].lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        root = ET.fromstring(archive.read(target))
        headers = None
        result = []
        for row in root.find(NS + "sheetData"):
            values = {}
            for cell in row:
                ref = cell.get("r")
                if not ref:
                    continue
                value = cell.find(NS + "v")
                inline = cell.find(NS + "is")
                if cell.get("t") == "s" and value is not None:
                    content = shared[int(value.text)]
                elif cell.get("t") == "inlineStr" and inline is not None:
                    content = "".join(t.text or "" for t in inline.iter(NS + "t"))
                else:
                    content = value.text if value is not None else ""
                values[_column_index(ref)] = content
            if headers is None:
                headers = values
                continue
            record = {name: values.get(index, "") for index, name in headers.items()}
            if record.get("公证录像位置名称") != "MAH00002":
                continue
            result.append({
                "row_id": str(record["总序号（17595）"]), "video": "MAH00002",
                "bv": record.get("BV号", ""),
                "title": record.get("侵权视频标题", ""),
                "uploader": record.get("上传者昵称", ""),
                "duration": duration_seconds(record.get("时长")),
                "views_claimed": "", "part_claimed": "",
                "video_type": record.get("视频性质", ""),
                "source_workbook": str(path), "source_sheet": "取证日期排序（总表）",
                "source_row": int(row.get("r") or 0),
            })
    if len(result) != 65:
        raise ValueError(f"expected 65 逐玉 rows, got {len(result)}")
    return result
