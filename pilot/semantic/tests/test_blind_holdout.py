"""WO-007: target human cells cannot enter machine inputs before blind scoring."""
import json
import tempfile
import unittest
from types import SimpleNamespace
from unittest import mock
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from pilot.ledger.workbook import duration_seconds
from pilot.ledger.pipeline import screenshot_ocr
from pilot.robustness import freeze, generate
from pilot.review.data import read_source
from pilot.semantic.table import load_table


class BlindHoldoutTest(unittest.TestCase):
    def test_target_output_hidden_but_outside_style_example_available(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "比对表"
            sheet.append(["总编号", *[f"input{i}" for i in range(2, 21)], *[f"output{i}" for i in range(21, 37)]])
            sheet.append([469, *(["input"] * 19), *(["held-out human"] * 16)])
            sheet.append([851, *(["example"] * 19), *(["style example"] * 16)])
            book.save(path)
            table = load_table(path, holdout_rows={469}, sample_rows={851})
            self.assertIsNone(table.value(469, "AF"))
            self.assertEqual(table.value(851, "AF"), "style example")
            _, source = read_source(path, "比对表", {2}, 36, {2})
            self.assertEqual(source[2][20:36], [""] * 16)

    def test_style_row_inside_holdout_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "比对表"
            sheet.append(["总编号", *[f"input{i}" for i in range(2, 21)], *[f"output{i}" for i in range(21, 37)]])
            sheet.append([469, *(["input"] * 19), *(["secret"] * 16)])
            book.save(path)
            with self.assertRaisesRegex(ValueError, "overlaps"):
                load_table(path, holdout_rows={469}, sample_rows={469})

    def test_mixed_download_directory_ocr_only_images(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "Screenshot_20250826_140316_phone.jpg"
            video = root / "Screenshot_20250826_140317_download.mp4"
            image.write_bytes(b"image")
            video.write_bytes(b"video")
            with mock.patch("pilot.ledger.pipeline.parallel_ocr", return_value=[{"frame": image.name}] ) as ocr:
                result = screenshot_ocr([image, video], "qsjx", root / "cache",
                                        r"Screenshot_\d{8}_(\d\d)(\d\d)(\d\d)_")
            self.assertEqual([r["path"] for r in result], [str(image)])
            ocr.assert_called_once_with([image])

    def test_blind_materials_require_frozen_machine_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = Workbook(); sheet = book.active; sheet.title = "比对表"
            sheet.append(["总编号", *[f"input{i}" for i in range(2, 21)],
                          *[f"output{i}" for i in range(21, 37)]])
            output = ["人工事实"] * 16; output[11] = "人工意见"
            sheet.append([469, *(["input"] * 19), *output])
            path = root / "source.xlsx"; book.save(path)
            case = root / "case"; vlm = case / "vlm"; rid = "469"
            for sub in ("content", "semantic", "robustness/deepseek", "vlm/r3/469"):
                (case / sub).mkdir(parents=True, exist_ok=True)
            (case / "local_stage_complete.json").write_text("{}")
            (case / "ledger.jsonl").write_text(json.dumps({"row_id": rid, "plaintiff":
                {"source_row": 2, "title": "标题"}}) + "\n")
            (case / "content/469.json").write_text("{}")
            (vlm / "r3/469/consensus.json").write_text('{"status": "ok"}')
            (case / "semantic/469.json").write_text(json.dumps({"opinion": {"text": "GLM意见"},
                "fields": {c: {"value": "机器事实"} for c in
                           "U V W X Y Z AA AB AC AD AE AG AH AI AJ".split()}}))
            (case / "robustness/deepseek/469.json").write_text(json.dumps({"text": "DeepSeek意见"}))
            config = SimpleNamespace(case_id="test_case", case_dir=case, vlm_dir=vlm, row_ids=[rid],
                output_indices=tuple(range(21, 37)), raw={"workbook": {"blind_holdout": True,
                    "path": str(path), "sheet": "比对表"}}, path_of=lambda value: Path(value))
            with self.assertRaises(FileNotFoundError):
                generate(config)
            freeze(config)
            (case / "content/469.json").write_text('{"changed": true}')
            with self.assertRaisesRegex(RuntimeError, "changed after freeze"):
                generate(config)
            freeze(config)
            result = generate(config)
            self.assertEqual(result["blind_opinion_rows"], 1)
            key = json.loads((case / "robustness/key.json").read_text())
            self.assertEqual(set(key["pairs"][rid].values()), {"人工", "GLM", "DeepSeek"})

    def test_excel_duration_past_24_encoded_minutes(self):
        self.assertEqual(duration_seconds(datetime(1900, 1, 1, 14, 3), qsjx=True), 2283)


if __name__ == "__main__":
    unittest.main()
