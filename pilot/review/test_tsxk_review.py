"""WO-104: synthetic G-group workbook and the five permitted sample recordings."""
from __future__ import annotations

import hashlib
import json
import threading
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
import yaml
from openpyxl.comments import Comment

from pilot.review.tsxk import OUTPUT, TsxkReviewStore, flag_columns, parse_ac
from pilot.review.tsxk_export import export_tsxk
from pilot.tsxk.config import TsxkCase


def fixture(tmp_path):
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    other = book.active
    other.title = "其他表"
    other["C3"] = "不要修改"
    other["C3"].comment = Comment("原有批注", "原作者")
    sheet = book.create_sheet("沈瑞曦")
    for col in range(1, 30):
        sheet.cell(1, col, openpyxl.utils.get_column_letter(col) + " 标题")
    sheet["B2"] = 7813
    sheet["B3"] = 7814
    sheet["B4"] = 9000
    sheet["C2"] = "自造视频一"
    sheet["C3"] = "自造视频二"
    sheet["E2"] = "https://www.bilibili.com/video/BVtest"
    sheet["Q4"] = "别的行原值"
    sheet["F2"] = "原表待替换"
    book.save(source)
    data = tmp_path / "state"
    (data / "rows").mkdir(parents=True)
    (data / "webledger").mkdir()
    (data / "timeline" / "7813" / "frames").mkdir(parents=True)
    (data / "audio").mkdir()
    evidence = tmp_path / "evidence"
    evidence.mkdir()
    (evidence / "7813.mp4").write_bytes(b"synthetic-video-bytes")
    for rid in ("7813", "7814"):
        result = {"fields": {col: {"value": f"机器 {rid} {col}", "by": "test", "basis": "自造依据"}
                             for col in OUTPUT},
                  "flags": [{"level": "red" if rid == "7813" else "yellow",
                             "reason": "Q：自造画面证据待复核" if rid == "7813" else "T：自造类型待复核"}],
                  "evidence": []}
        result["fields"]["AC"]["value"] = "\n\n".join(
            f"{heading}\n{rid} 的 {key}" for key, heading in
            (("主题概括", "【主题概括】"), ("段落情节", "【段落情节】"),
             ("新增创作要素", "【新增创作要素情况】"), ("比对意见", "【比对意见】"),
             ("其他", "【其他】")))
        (data / "rows" / f"{rid}.json").write_text(json.dumps(result, ensure_ascii=False))
        (data / "webledger" / f"{rid}.json").write_text(json.dumps({
            "title": f"自造视频{rid}", "play_start": 1, "play_end": 8,
            "content_start": 0, "content_end": 7, "complete": rid == "7813",
            "complete_reason": "自造读数", "other_pages": [{"start": 8, "title": "自动连播"}]}))
    (data / "timeline" / "7813" / "timeline.json").write_text(json.dumps({
        "segments": [{"start": 1, "end": 8, "content": "tsxk", "form": "", "sources": ["自造"]}],
        "frames": [{"t": 1, "content": "tsxk"}], "summary": {}, "flags": []}))
    (data / "timeline" / "7813" / "frames" / "00001.00.jpg").write_bytes(b"test-image")
    (data / "audio" / "7813.json").write_text(json.dumps({
        "status": "ok", "played_seconds": 7,
        "segments": [{"kind": "uploader", "text": "自造口播", "recording_start": 2, "recording_end": 4}],
        "promotional_speech": [{"text": "自造推广", "recording_start": 5, "recording_end": 6}]}))
    config = tmp_path / "case.yaml"
    config.write_text(yaml.safe_dump({
        "schema": "tsxk-web-1", "case_id": "tsxk-test",
        "workbook": {"path": str(source), "sheet": "沈瑞曦", "rows": {"first": 7813, "last": 7814},
                     "id_column": "B", "input_columns": ["B", "C", "E"],
                     "output_columns": list(OUTPUT)},
        "evidence": {"dir": str(evidence), "pattern": "{row}.mp4"},
        "webledger": {"dir": str(data / "webledger")},
        "output": {"dir": str(data)}, "rules": str(tmp_path / "unused.yaml")
    }, allow_unicode=True))
    return source, config, data


def test_review_decisions_and_precise_export(tmp_path):
    source, config, data = fixture(tmp_path)
    original_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    store = TsxkReviewStore(TsxkCase(config))
    assert store.excel_rows == {"7813": 2, "7814": 3}
    detail = store.detail("7813")
    assert detail["video_url"] == "/api/video/7813"
    assert detail["timeline"]["segments"][0]["thumbnail"].endswith("00001.00.jpg")
    assert detail["sound"]["uploader_speech"][0]["text"] == "自造口播"
    assert len(next(f for f in detail["fields"] if f["column"] == "AC")["sections"]) == 5
    store.decide("7813", "approve", {"D": "机器 7813 D"}, actor="Owner A")
    store.decide("7813", "modify", {"F": "所有者修订的网址"}, actor="Owner B")
    store.decide("7813", "approve", {"F": "所有者修订的网址"}, actor="Owner B")
    reopened = TsxkReviewStore(TsxkCase(config))
    assert reopened.summary("7813")["reviewed"] == 2
    assert next(f for f in reopened.fields("7813") if f["column"] == "F")["modified"]
    assert [event["who"] for event in reopened.events["7813"]] == ["Owner A", "Owner B", "Owner B"]
    dest = export_tsxk(reopened, tmp_path / "copy.xlsx")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_sha
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(dest) as b:
        assert a.read("xl/worksheets/sheet1.xml") == b.read("xl/worksheets/sheet1.xml")
        assert a.read("xl/workbook.xml") == b.read("xl/workbook.xml")
    original = openpyxl.load_workbook(source)
    copied = openpyxl.load_workbook(dest)
    assert copied.sheetnames == ["其他表", "沈瑞曦"]
    assert copied["其他表"]["C3"].value == "不要修改"
    assert copied["其他表"]["C3"].comment.text == "原有批注"
    assert copied["沈瑞曦"]["D2"].value == "机器 7813 D"
    assert copied["沈瑞曦"]["F2"].value == "所有者修订的网址"
    assert copied["沈瑞曦"]["Q2"].comment.text.startswith("标红：Q")
    assert copied["沈瑞曦"]["Q2"].fill.fgColor.rgb == "FFF7C8C3"
    assert copied["沈瑞曦"]["T3"].comment.text.startswith("标黄：T")
    assert copied["沈瑞曦"]["Q4"].value == "别的行原值"
    allowed = {f"{col}{row}" for col in OUTPUT for row in (2, 3)}
    for row in original["沈瑞曦"]:
        for cell in row:
            if cell.coordinate not in allowed:
                other = copied["沈瑞曦"][cell.coordinate]
                assert (other.value, other.style_id, bool(other.comment)) == (
                    cell.value, cell.style_id, bool(cell.comment))
    original.close()
    copied.close()
    report = dest.with_suffix(".md").read_text()
    assert "7813：F" not in report  # F was reviewed
    assert "7814：D、F、L" in report
    assert "Owner B" in report and "所有者修订的网址" in report


def test_missing_machine_row_blocks_export(tmp_path):
    _, config, data = fixture(tmp_path)
    (data / "rows" / "7814.json").unlink()
    with pytest.raises(ValueError, match="7814"):
        export_tsxk(TsxkReviewStore(TsxkCase(config)), tmp_path / "no.xlsx")
    assert not (tmp_path / "no.xlsx").exists()


def test_equals_sign_flag_targets_its_column():
    assert flag_columns({"reason": "Q=是：判为完全不含动画，交人判断"}) == ["Q"]


def test_case_commands_dispatch_to_tsxk(tmp_path, capsys):
    _, config, data = fixture(tmp_path)
    from pilot.__main__ import main
    with patch("pilot.review.server.serve") as serve:
        main(["review", str(config)])
    assert serve.call_args.kwargs["store"].kind == "tsxk"
    main(["export", str(config)])
    result = capsys.readouterr().out.strip()
    assert result.endswith(".xlsx")
    assert Path(result).is_file()
    assert Path(result).parent == data / "export"


def test_loopback_review_and_video_range(tmp_path, monkeypatch):
    _, config, _ = fixture(tmp_path)
    store = TsxkReviewStore(TsxkCase(config))
    from http.server import ThreadingHTTPServer
    from pilot.review import server as module
    servers = []
    def capture(address, handler):
        servers.append(ThreadingHTTPServer(address, handler))
        return servers[-1]
    monkeypatch.setattr(module, "ThreadingHTTPServer", capture)
    thread = threading.Thread(target=module.serve, args=(store.case,),
                              kwargs={"port": 0, "store": store}, daemon=True)
    thread.start()
    for _ in range(100):
        if servers:
            break
        time.sleep(.01)
    assert servers and servers[0].server_address[0] == "127.0.0.1"
    base = f"http://127.0.0.1:{servers[0].server_address[1]}"
    try:
        with urllib.request.urlopen(base + "/", timeout=5) as response:
            assert "吞噬星空" in response.read().decode()
        with urllib.request.urlopen(base + "/api/row/7813", timeout=5) as response:
            assert json.load(response)["row_id"] == "7813"
        request = urllib.request.Request(base + "/api/video/7813", headers={"Range": "bytes=0-3"})
        with urllib.request.urlopen(request, timeout=5) as response:
            assert response.status == 206
            assert response.read() == b"synt"
        with pytest.raises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(base + "/api/video/9000", timeout=5)
        assert error.value.code == 404
        foreign = urllib.request.Request(base + "/api/rows", headers={"Origin": "https://other.example"})
        with pytest.raises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(foreign, timeout=5)
        assert error.value.code == 403
    finally:
        servers[0].shutdown()
        thread.join(timeout=3)


@pytest.mark.parametrize("sample", range(13, 18))
def test_sample_output_has_five_ac_sections(sample):
    path = Path(__file__).resolve().parents[2] / "workspace" / "tsxk" / "samples" / str(sample) / "fields.json"
    if not path.exists():
        pytest.skip("local sample outputs are not installed")
    result = json.loads(path.read_text(encoding="utf-8"))
    assert set(OUTPUT) <= set(result["fields"])
    sections = parse_ac(result["fields"]["AC"]["value"])
    assert len(sections) == 5
    assert all(section["heading"] for section in sections)
