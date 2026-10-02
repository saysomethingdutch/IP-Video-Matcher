"""Cross-module check: logical fields land in the configured physical Excel columns."""
from __future__ import annotations

import json

import openpyxl
import yaml

from pilot.case_config import CaseConfig, LOGICAL_OUTPUTS
from pilot.review.data import ReviewStore
from pilot.review.export import export_case
from pilot.semantic.table import load_table


def test_review_and_export_respect_case_column_mapping(tmp_path):
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Case Sheet"
    for col in range(1, 53):
        sheet.cell(1, col, f"Header {col}")
    sheet.cell(2, 1, 17)
    sheet.cell(2, 21, "original U")
    book.save(source)

    mapping = {letter: openpyxl.utils.get_column_letter(37 + i)
               for i, letter in enumerate(LOGICAL_OUTPUTS)}
    output = tmp_path / "case"
    config_path = tmp_path / "case.yaml"
    rules = tmp_path / "opinion.txt"
    rules.write_text("Write from the facts only.")
    config_path.write_text(yaml.safe_dump({
        "schema_version": 1, "case_id": "test_case",
        "workbook": {"path": str(source), "sheet": "Case Sheet", "row_ids": [17],
                     "columns": {"row_id": "Header 1", "video": "Header 2"}, "bv_column": "L",
                     "output_columns": mapping},
        "evidence": {"videos": {"V1": {}}, "video_dir": str(tmp_path), "screenshot_dir": str(tmp_path)},
        "source": {"enabled": False},
        "client": {"rules": str(tmp_path / "rules.yaml"), "work_title": "Test Work",
                   "style_examples": {}, "opinion_format": {
            "headings": ["Purpose", "Proportion", "Substitution"], "conclusion": "End.",
            "prompt_rules": str(rules)}},
        "models": {"visual": {"allow_image_upload": False}, "text": "local", "speech": {"enabled": False}},
        "output": {"dir": str(output), "export_name": "result"}, "ledger": {"case": "custom", "screenshot_time_pattern": r"(\d\d)(\d\d)(\d\d)"}
    }), encoding="utf-8")
    config = CaseConfig(config_path)
    table = load_table(source, config.raw["workbook"]["sheet"], mapping, "L")
    assert table.header("U") == "Header 37"
    output.mkdir()
    ledger = {"row_id": "17", "video": "V1", "video_duration": 5,
              "plaintiff": {"source_row": 2, "source_workbook": str(source), "uploader": "u", "title": "t"}}
    (output / "ledger.jsonl").write_text(json.dumps(ledger) + "\n")
    semantic = output / "semantic"
    semantic.mkdir()
    (semantic / "17.json").write_text(json.dumps({"fields": {
        letter: {"value": f"machine {letter}", "basis": "test"} for letter in LOGICAL_OUTPUTS},
        "flags": [{"level": "yellow", "reason": "无 ASR（D-020 中止）"}]}))
    audio = output / "audio"
    audio.mkdir()
    (audio / "17.json").write_text(json.dumps({"status": "ok", "engine": "SenseVoice-Small", "segments": []}))

    store = ReviewStore(output, config)
    assert store.detail("17")["fields"][36]["value"] == "machine U"
    assert store.detail("17")["audio"]["status"] == "ok"
    assert all("无 ASR" not in flag["reason"] for flag in store.detail("17")["flags"])
    exported = export_case(store)
    filled = openpyxl.load_workbook(exported, read_only=True, data_only=True)
    assert filled["Case Sheet"]["U2"].value == "original U"
    assert filled["Case Sheet"]["AK2"].value == "machine U"
    assert filled["Case Sheet"]["AZ2"].value == "machine AJ"
    assert "机器来源" in filled.sheetnames
    filled.close()

    # The reusable delegated-list adapter uses physical column names, including
    # remapped cases, and records the delegate separately from the Owner.
    from pilot.review.batch import preview, apply
    items=[{'row':'17','column':'AK','old':'machine U','new':'approved duration','reason':'checked by Owner'}]
    proposal=preview(store,items)
    assert proposal['matched']==1
    answer=apply(store,items,delegate='assistant',confirmed=True,preview_token=proposal['token'])
    assert answer['applied']==1
    assert store.detail('17')['fields'][36]['value']=='approved duration'
    event=json.loads(store.decisions_path.read_text().splitlines()[-1])
    assert event['field']=='37' and event['who']=='Owner 代办'
    assert event['delegate']=='assistant' and event['reason']=='checked by Owner'
    assert apply(store,items,delegate='assistant',confirmed=True,preview_token=proposal['token'])['applied']==0
