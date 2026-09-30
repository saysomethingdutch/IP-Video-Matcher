"""Review data and append-only owner decisions for the 吞噬星空 workbook."""
from __future__ import annotations

import csv
import json
import os
import re
import threading
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from pilot.tsxk.ac import HEADINGS
from pilot.tsxk.config import TsxkCase

OUTPUT = ("D", "F", "L", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC")
CLASSIFICATION = OUTPUT[3:]
LEVELS = {"yellow": 1, "red": 2}
FIELD_MENTION = re.compile(r"(?<![A-Za-z0-9])(?:AA|AB|AC|D|F|L|Q|R|S|T|U|V|W|X|Y|Z)(?=[:：=＝、，\s]|$)")
AC_HEADING = re.compile(r"【(?:主题概括|段落情节|新增创作要素情况|比对意见|其他)】")
FOCUS_REASON = re.compile(r"复核后两家仍有分歧|画面未能判定|画面仍待定|争议帧.*待定|分类.*待定|R[：:=＝].*(?:待核|复核)|完整播放.*(?:待核|复核)")


def parse_ac(value: str) -> list[dict]:
    """Keep the five editable sections visible even when a machine section is empty."""
    text = str(value or "")
    found = list(AC_HEADING.finditer(text))
    parts = {}
    for i, match in enumerate(found):
        end = found[i + 1].start() if i + 1 < len(found) else len(text)
        parts[match.group()] = text[match.end():end].strip()
    if not found and text.strip():
        parts[HEADINGS[0][1]] = text.strip()
    return [{"key": key, "heading": heading, "text": parts.get(heading, "")}
            for key, heading in HEADINGS]


def flag_columns(flag: dict) -> list[str]:
    explicit = flag.get("columns") or flag.get("column")
    if isinstance(explicit, str):
        explicit = [explicit]
    if explicit:
        return [col for col in OUTPUT if col in explicit]
    reason = str(flag.get("reason") or "")
    mentioned = set(FIELD_MENTION.findall(reason))
    if mentioned:
        return [col for col in OUTPUT if col in mentioned]
    if reason.startswith("分类") or "时间线" in reason:
        return list(CLASSIFICATION)
    return list(OUTPUT)


def normalize_flags(machine: dict | None, ledger: dict | None,
                    audit: list[dict] | None = None) -> list[dict]:
    raw = list((machine or {}).get("flags") or [])
    if not machine:
        raw.append({"level": "red", "reason": "正式行分类输出尚未完成", "columns": list(OUTPUT)})
    if (ledger or {}).get("complete_needs_review"):
        raw.append({"level": "yellow", "reason": "R：台账完整播放判断需要人工复核", "columns": ["R"]})
    values = {col: str(((machine or {}).get("fields") or {}).get(col, {}).get("value") or "").strip()
              for col in ("Q", "R", "Y", "Z")}
    for col, value, reason in (("Z", "高", "Z=高：合理使用可能性需人工确认"),
                               ("Q", "是", "Q=是：请核对完全不含动画的判断")):
        if values[col] == value and not any(col in flag_columns(x) and
                (str(x.get("level", "")).lower() in ("red", "error") or
                 (col == "Z" and "Z 高" in str(x.get("reason", ""))))
                for x in raw if isinstance(x, dict)):
            raw.append({"level": "yellow", "reason": reason, "columns": [col], "tier": "focus"})
    if "偷渡" in values["Y"]:
        raw.append({"level": "yellow", "reason": "Y=偷渡：请核对无关内容及时间段", "columns": ["Y"], "tier": "focus"})
    if not values["R"] and not (ledger or {}).get("complete_needs_review"):
        raw.append({"level": "yellow", "reason": "R 待核：完整播放尚不能确定", "columns": ["R"], "tier": "focus"})
    for finding in audit or []:
        columns = [col.strip() for col in str(finding.get("columns") or "").split(",")]
        reason = f"来自独立审计（{finding.get('type', '')}）：{finding.get('explanation', '')}"
        if finding.get("evidence"):
            reason += f"；线索：{finding['evidence']}"
        raw.append({"level": "yellow", "reason": reason,
                    "columns": [col for col in OUTPUT if col in columns], "tier": "focus",
                    "source": "audit"})
    seen = set()
    result = []
    for item in raw:
        item = item if isinstance(item, dict) else {"level": "yellow", "reason": str(item)}
        reason = str(item.get("reason") or item.get("message") or "").strip()
        level = "red" if str(item.get("level") or "").lower() in ("red", "error") else "yellow"
        columns = flag_columns(item)
        tier = "focus" if level == "red" or item.get("tier") != "technical" and (item.get("tier") == "focus" or FOCUS_REASON.search(reason) or (
            "Z 高" in reason or "Q=是" in reason)) else "technical"
        key = (level, reason, tuple(columns))
        if not reason or key in seen:
            continue
        seen.add(key)
        result.append({"level": level, "reason": reason, "columns": columns,
                       "tier": tier, "source": item.get("source", "machine")})
    return result


class TsxkReviewStore:
    kind = "tsxk"

    def __init__(self, case: TsxkCase, decisions_path: Path | None = None):
        self.config = case
        self.case = case.case_dir
        self.row_ids = case.rows
        self.sheet_name = case.raw["workbook"]["sheet"]
        if tuple(case.raw["workbook"]["output_columns"]) != OUTPUT:
            raise ValueError("configured output columns differ from WO-104")
        if case.case_id == "tsxk" and (
            self.sheet_name != "沈瑞曦" or self.row_ids != [str(n) for n in range(7813, 8213)]
        ):
            raise ValueError("WO-104 may only review 沈瑞曦 rows 7813–8212")
        self.source_path = case.path_of(case.raw["workbook"]["path"])
        self.output_columns = tuple(openpyxl.utils.column_index_from_string(col) for col in OUTPUT)
        self.export_name = "【G组】吞噬星空视频比对-机器核对"
        self.decisions_path = Path(decisions_path) if decisions_path is not None else self.case / "review" / "decisions.jsonl"
        self.audit: dict[str, list[dict]] = defaultdict(list)
        audit_path = self.case / "audit" / "findings.csv"
        if audit_path.is_file():
            with audit_path.open(encoding="utf-8-sig", newline="") as handle:
                for finding in csv.DictReader(handle):
                    if str(finding.get("row") or "") in case.rows:
                        self.audit[str(finding["row"])].append(finding)
        self.lock = threading.RLock()
        self.latest: dict[str, dict[str, dict]] = defaultdict(dict)
        self.events: dict[str, list[dict]] = defaultdict(list)
        wanted = set(self.row_ids)
        id_col = openpyxl.utils.column_index_from_string(case.raw["workbook"]["id_column"])
        book = openpyxl.load_workbook(self.source_path, read_only=True, data_only=True)
        try:
            sheet = book[self.sheet_name]
            self.headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            self.excel_rows = {}
            for row_number, cells in enumerate(sheet.iter_rows(min_row=2, min_col=id_col,
                                                                 max_col=id_col, values_only=True), 2):
                value = cells[0]
                if value is None:
                    continue
                rid = str(int(value)) if isinstance(value, (int, float)) and float(value).is_integer() else str(value).strip()
                if rid in wanted:
                    if rid in self.excel_rows:
                        raise ValueError(f"duplicate row ID in {self.sheet_name}: {rid}")
                    self.excel_rows[rid] = row_number
        finally:
            book.close()
        missing = wanted - self.excel_rows.keys()
        if missing:
            raise ValueError(f"{self.sheet_name} is missing {len(missing)} configured IDs: {sorted(missing)[:5]}")
        self.inputs = case.read_inputs(rows=self.row_ids)
        if self.decisions_path.exists():
            for line in self.decisions_path.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    event = json.loads(line)
                    rid, col = str(event.get("row")), str(event.get("field"))
                    if rid in wanted and col in OUTPUT:
                        self.latest[rid][col] = event
                        self.events[rid].append(event)

    def _load(self, path: Path) -> dict | None:
        if not path.is_file():
            return None
        return json.loads(path.read_text(encoding="utf-8"))

    def machine(self, rid: str) -> dict | None:
        return self._load(self.case / "rows" / f"{rid}.json")

    def ledger(self, rid: str) -> dict | None:
        return self.config.webledger(rid)

    def flags(self, rid: str, machine: dict | None = None, ledger: dict | None = None) -> list[dict]:
        return normalize_flags(machine if machine is not None else self.machine(rid),
                               ledger if ledger is not None else self.ledger(rid), self.audit.get(rid))

    def fields(self, rid: str, machine: dict | None = None, flags: list[dict] | None = None) -> list[dict]:
        machine = machine if machine is not None else self.machine(rid)
        flags = flags if flags is not None else self.flags(rid, machine)
        source = (machine or {}).get("fields") or {}
        result = []
        for col in OUTPUT:
            raw = source.get(col) or {}
            if not isinstance(raw, dict):
                raw = {"value": raw}
            value = raw.get("value")
            machine_value = "" if value is None else str(value)
            event = self.latest[rid].get(col)
            final = str(event["new"]) if event else machine_value
            related = [flag for flag in flags if col in flag["columns"]]
            focus = [flag for flag in related if flag["tier"] == "focus"]
            technical = [flag for flag in related if flag["tier"] == "technical"]
            level = max((flag["level"] for flag in focus), key=lambda x: LEVELS[x], default=None)
            index = openpyxl.utils.column_index_from_string(col) - 1
            modified = any(e["field"] == col and e["action"] == "modify" for e in self.events[rid])
            options = self.allowed_values(col)
            invalid = bool(options is not None and final not in options)
            result.append({"column": col, "name": str(self.headers[index] or col),
                           "machine_value": machine_value, "value": final,
                           "by": raw.get("by") or "", "basis": raw.get("basis") or "",
                           "reviewed": bool(event), "modified": modified,
                           "reviewer": event.get("who") if event else None,
                           "reviewed_at": event.get("when") if event else None,
                           "retain_mark": bool(event.get("retain_mark")) if event else False,
                           "mark_note": str(event.get("mark_note") or "") if event else "",
                           "flag_level": level, "flag_reasons": [x["reason"] for x in focus],
                           "technical_reasons": [x["reason"] for x in technical],
                           "flag_details": [{"level": x["level"], "reason": x["reason"]} for x in focus],
                           "sections": parse_ac(final) if col == "AC" else None,
                           "options": options, "invalid_value": invalid})
        return result

    def allowed_values(self, col: str) -> list[str] | None:
        if col not in ("Q", "R", "T", "U", "V", "W", "X", "Y", "Z"):
            return None
        if col == "T":
            from itertools import combinations
            forms = ("漫画", "手游", "小说", "有声书")
            return ["否"] + [f"是（{'、'.join(group)}）" for n in range(1, len(forms)+1) for group in combinations(forms, n)]
        if not hasattr(self, "_allowed_rules"):
            self._allowed_rules = self.config.rules()["columns"]
        rules = self._allowed_rules
        return list(rules[col]["values"])

    def cleanup_preview(self) -> dict:
        """Read-only proposal. The manifest token makes a stale confirmation fail."""
        import hashlib
        from pilot.tsxk.ac_filter import cleanup_other
        result = []
        with self.lock:
            for rid in self.row_ids:
                event = self.latest[rid].get("AC")
                if not event:
                    continue
                current = str(event["new"])
                replacement, removed = cleanup_other(current)
                if removed:
                    result.append({"row": rid, "old": current, "new": replacement, "removed": removed})
        token = hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        return {"rows": result, "row_count": len(result), "sentence_count": sum(len(r['removed']) for r in result), "token": token}

    def cleanup_confirm(self, token: str, actor: str, confirmed: bool) -> dict:
        if confirmed is not True:
            raise ValueError("请先查看清单，再明确确认清理")
        actor = str(actor or "").strip()
        if not actor or actor == "WO104 临时试跑" or len(actor) > 80:
            raise ValueError("请填写自己的核验人名字")
        with self.lock:
            proposal = self.cleanup_preview()
            if token != proposal["token"]:
                raise ValueError("清单已变化，请重新查看后确认")
            events = []
            for row in proposal["rows"]:
                current = self.latest[row['row']]['AC']
                events.append({"when": datetime.now(timezone.utc).isoformat(), "who": actor, "row": row['row'],
                               "field": "AC", "action": "modify", "old": row['old'], "new": row['new'],
                               "machine": current.get('machine', ''), "retain_mark": current.get('retain_mark', False),
                               "mark_note": current.get('mark_note', ''), "reason": "所有者确认：清理【其他】中的模型说明"})
            if events:
                self._write_events(events)
            return {"rows": len(events), "sentences": proposal['sentence_count']}

    def summary(self, rid: str) -> dict:
        if rid not in self.excel_rows:
            raise KeyError(rid)
        machine = self.machine(rid)
        ledger = self.ledger(rid)
        flags = self.flags(rid, machine, ledger)
        y = self.latest[rid].get("Y")
        category = str(y["new"]) if y else str(((machine or {}).get("fields") or {}).get("Y", {}).get("value") or "待分类")
        reviewed = len(self.latest[rid])
        focus = [flag for flag in flags if flag["tier"] == "focus"]
        unresolved = any(col not in self.latest[rid] for flag in focus for col in flag["columns"])
        return {"row_id": rid, "title": str(self.inputs.get(rid, {}).get("C") or (ledger or {}).get("title") or ""),
                "category": category, "status": "待核" if reviewed < len(OUTPUT) else
                "已修改" if any(e["action"] == "modify" for e in self.events[rid]) else "已通过",
                "reviewed": reviewed, "total": len(OUTPUT), "machine_ready": machine is not None,
                "yellow": sum(f["level"] == "yellow" for f in focus),
                "red": sum(f["level"] == "red" for f in focus),
                "technical": sum(f["tier"] == "technical" for f in flags),
                "priority_unreviewed": unresolved}

    def all_summaries(self) -> list[dict]:
        return [self.summary(rid) for rid in self.row_ids]

    def _timeline(self, rid: str) -> dict | None:
        return self._load(self.case / "timeline" / rid / "timeline.json")

    def thumbnail_path(self, rid: str, name: str) -> Path | None:
        if rid not in self.excel_rows or Path(name).name != name or not name.lower().endswith(".jpg"):
            return None
        root = (self.case / "timeline" / rid / "frames").resolve()
        path = (root / name).resolve()
        return path if path.is_relative_to(root) and path.is_file() else None

    def video_path(self, rid: str) -> Path | None:
        if rid not in self.excel_rows:
            return None
        video = self.config.evidence(rid)
        return video if video.is_file() else None

    def detail(self, rid: str) -> dict:
        from .tsxk_views import view_count_detail
        if rid not in self.excel_rows:
            raise KeyError(rid)
        machine = self.machine(rid)
        ledger = self.ledger(rid) or {}
        flags = self.flags(rid, machine, ledger)
        timeline = self._timeline(rid) or {}
        sound = self._load(self.case / "audio" / f"{rid}.json") or {}
        frames = timeline.get("frames") or []
        segments = []
        for segment in timeline.get("segments") or []:
            match = next((f for f in frames if segment.get("start", 0) <= f.get("t", -1) < segment.get("end", 0)), None)
            name = f"{float(match['t']):08.2f}.jpg" if match else ""
            thumb = f"/api/thumb/{rid}/{name}" if name and self.thumbnail_path(rid, name) else None
            segments.append({k: segment.get(k) for k in ("start", "end", "content", "form", "sources")}
                            | {"thumbnail": thumb})
        speech = [s for s in sound.get("segments") or [] if s.get("kind") == "uploader"]
        return {**self.summary(rid), "fields": self.fields(rid, machine, flags),
                "view_count": view_count_detail(self.config, rid),
                "flags": flags, "history": self.events[rid], "excel_row": self.excel_rows[rid],
                "bilibili_url": self.inputs.get(rid, {}).get("E") or ledger.get("address"),
                "video_url": f"/api/video/{rid}" if self.video_path(rid) else None,
                "ledger": {k: ledger.get(k) for k in ("play_start", "play_end", "content_start",
                    "content_end", "content_basis", "complete", "complete_basis", "complete_reason",
                    "complete_needs_review", "passes", "other_pages", "flags")},
                "timeline": {"segments": segments, "flags": timeline.get("flags") or [],
                             "summary": timeline.get("summary") or {}},
                "sound": {"status": sound.get("status"), "uploader_speech": speech,
                          "uploader_text": sound.get("uploader_speech") or [],
                          "promotional_speech": sound.get("promotional_speech") or [],
                          "played_seconds": sound.get("played_seconds")},
                "evidence": (machine or {}).get("evidence") or []}

    def _write_events(self, events: list[dict]) -> None:
        self.decisions_path.parent.mkdir(parents=True, exist_ok=True)
        with self.decisions_path.open("a", encoding="utf-8") as handle:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
            handle.write("".join(json.dumps(event, ensure_ascii=False) + "\n" for event in events))
            handle.flush()
            os.fsync(handle.fileno())
        for event in events:
            rid, col = event["row"], event["field"]
            self.latest[rid][col] = event
            self.events[rid].append(event)

    def decide(self, rid: str, action: str, edits: dict | None = None, actor: str | None = None,
               retain_mark: bool = False, mark_note: str = "") -> dict:
        if rid not in self.excel_rows or action not in ("approve", "modify", "mark"):
            raise ValueError("invalid row or action")
        actor = str(actor or "").strip()
        if not 1 <= len(actor) <= 80:
            raise ValueError("请填写核验人（1–80 字）")
        edits = edits or {}
        if len(edits) != 1 or next(iter(edits)) not in OUTPUT:
            raise ValueError("每次只能核验一个产出列")
        col, requested = next(iter(edits.items()))
        if self.machine(rid) is None:
            raise ValueError("该行机器输出尚未生成")
        if type(retain_mark) is not bool:
            raise ValueError("导出标记必须是勾选状态")
        mark_note = str(mark_note or "").strip()
        if len(mark_note) > 1000 or retain_mark and not mark_note:
            raise ValueError("保留标记时请写所有者说明（最多 1000 字）")
        with self.lock:
            field = next(x for x in self.fields(rid) if x["column"] == col)
            old = field["value"]
            new = str(requested) if action == "modify" else old
            options = self.allowed_values(col)
            if action != "mark" and options is not None and new not in options:
                raise ValueError(f"{col} 当前值不合规，请从下拉框重选并保存")
            if action == "modify" and new == old:
                raise ValueError("修改值与当前值相同")
            if action == "mark" and not field["reviewed"]:
                raise ValueError("请先核验该格，再设置导出标记")
            if retain_mark and not field["flag_level"]:
                raise ValueError("此格没有重点标记可保留")
            if action == "mark" and retain_mark == field["retain_mark"] and mark_note == field["mark_note"]:
                raise ValueError("导出标记没有变化")
            event = {"when": datetime.now(timezone.utc).isoformat(), "who": actor,
                     "row": rid, "field": col, "action": action, "old": old, "new": new,
                     "machine": field["machine_value"], "retain_mark": retain_mark,
                     "mark_note": mark_note if retain_mark else ""}
            self._write_events([event])
            return self.summary(rid)

    def approve_row(self, rid: str, actor: str | None = None) -> dict:
        if rid not in self.excel_rows or self.machine(rid) is None:
            raise ValueError("该行机器输出尚未生成")
        actor = str(actor or "").strip()
        if not 1 <= len(actor) <= 80:
            raise ValueError("请填写核验人（1–80 字）")
        with self.lock:
            stamp = datetime.now(timezone.utc).isoformat()
            events = [{"when": stamp, "who": actor, "row": rid, "field": field["column"],
                       "action": "approve", "old": field["value"], "new": field["value"],
                       "machine": field["machine_value"], "retain_mark": False, "mark_note": ""}
                      for field in self.fields(rid) if not field["reviewed"]]
            if any(f["invalid_value"] for f in self.fields(rid) if not f["reviewed"]):
                raise ValueError("本行有不合规值，请从下拉框重选并保存后批准")
            if events:
                self._write_events(events)
            return {**self.summary(rid), "approved_now": len(events)}

    def export_row(self, rid: str) -> tuple[list[dict], list[dict]]:
        machine = self.machine(rid)
        if machine is None:
            raise ValueError(f"missing machine output: {rid}")
        source = machine.get("fields") or {}
        absent = [col for col in OUTPUT if col not in source]
        if absent:
            raise ValueError(f"machine output {rid} lacks {','.join(absent)}")
        flags = self.flags(rid, machine, self.ledger(rid))
        return self.fields(rid, machine, flags), flags

    def corrections(self) -> list[dict]:
        return [event for rid in self.row_ids for event in self.events[rid]
                if event.get("action") == "modify" and event.get("old") != event.get("new")]
